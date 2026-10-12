"""Offline safeguards: real cited NCSI values only; no invented market fixtures."""
from copy import deepcopy
from datetime import date
import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from src.manual_ingestion import (ROOT, coverage, decimal_value, deduplicate, digest,
    init_metadata, map_columns, normalize_rows, period, read_table, run, validate_metadata)
from src.manual_reference import GOVERNORATES, governorate
from src.process_data import parse_page


def genuine_national_observation():
    source=json.loads((ROOT/'docs/source_catalog.json').read_text())[0]
    excerpt=(ROOT/'tests/fixtures/ncsi_2025_10_activity_excerpt.txt').read_text()
    return next(r for r in parse_page(excerpt,source,'test-transcription-only')
                if r['indicator']=='sales_contract_count' and r['period_end']=='2025-09-30'
                and r['frequency']=='monthly')


def metadata(payload):
    # A temporary review declaration for the cited test transcription only,
    # never admitted as an original production workbook or committed dataset.
    return dict(dataset_id='cited_ncsi_test_transcription',title='NCSI October 2025 table 29 test transcription',
        publisher='NCSI',source_url='https://api.ncsi.gov.om/uploads/pdfs/publicationpdf1761108337.pdf',
        publication_date=None,retrieved_at='2026-10-12T00:00:00+04:00',expected_sha256=digest(payload),
        verification=dict(status='reviewed_official',reviewed_by='test_fixture_transcription',reviewed_at='2026-10-12T00:00:00+04:00'),
        licence=dict(name='Open Government Licence – Sultanate of Oman',url='https://www.ncsi.gov.om/assets/docs/gov_license.pdf',redistribution_allowed=True),
        tables=[table()])


def table():
    return dict(scope_id='ncsi_national_sales_contracts',observation_grain='aggregate',
        geographic_level='national',metrics=[dict(field='transaction_count',indicator='sales_contract_count',
            source_unit='contract count',unit='count')])


def actual_frame():
    row=genuine_national_observation()
    return pd.DataFrame([{'period_start':row['period_start'],'period_end':row['period_end'],
                         'frequency':row['frequency'],'transaction_count':row['value'],'__source_row':2}])


def normalized_actual_rows():
    return normalize_rows(actual_frame(),table(),metadata(b''),Path('cited-test-transcription.csv'),
                          'test-transcription-only',date(2026,10,12))[0]


@pytest.mark.parametrize('slug,english,arabic,aliases',GOVERNORATES)
def test_all_reference_governorates_bilingual(slug,english,arabic,aliases):
    assert governorate(english)[0]==slug
    assert governorate('محافظة '+arabic)[0]==slug
    for label in aliases:
        assert governorate(label)[0]==slug


def test_governorate_combined_regions_never_split():
    with pytest.raises(ValueError,match='Unknown or combined'):
        governorate('Al Batinah')
    assert governorate('الإجمالي') is None


def test_bilingual_header_mapping_with_no_records():
    frame=pd.DataFrame(columns=['السنة','المحافظة','عدد المعاملات','القيمة المتداولة'])
    assert map_columns(frame,{})=={'year':'السنة','governorate':'المحافظة',
                                  'transaction_count':'عدد المعاملات','transaction_value':'القيمة المتداولة'}


def test_explicit_mapping_and_unmapped_columns():
    frame=pd.DataFrame(columns=['Gov label','Contracts'])
    assert map_columns(frame,{'columns':{'governorate':'Gov label','transaction_count':'Contracts'}})['governorate']=='Gov label'
    with pytest.raises(ValueError,match='Unmapped columns'):
        map_columns(frame,{})


def test_ambiguous_headers_require_review():
    with pytest.raises(ValueError,match='Multiple headers'):
        map_columns(pd.DataFrame(columns=['Year','السنة']),{})


def test_locale_conversion_uses_actual_cited_value():
    actual=genuine_national_observation()['value']
    assert decimal_value(actual,'count')==actual
    arabic=actual.translate(str.maketrans('0123456789','٠١٢٣٤٥٦٧٨٩'))
    assert decimal_value(arabic,'count')==actual
    source=json.loads((ROOT/'docs/source_catalog.json').read_text())[0]
    excerpt=(ROOT/'tests/fixtures/ncsi_2025_10_activity_excerpt.txt').read_text()
    value=next(r['value'] for r in parse_page(excerpt,source,'test-only')
               if r['indicator']=='sales_value_omr' and r['frequency']=='monthly' and r['period_end']=='2025-09-30')
    millions=str(__import__('decimal').Decimal(value)/1000000)
    assert __import__('decimal').Decimal(decimal_value(millions,'OMR','1000000'))==__import__('decimal').Decimal(value)
    assert decimal_value(millions.replace('.',','),'OMR','1000000',',','.')==decimal_value(millions,'OMR','1000000')


@pytest.mark.parametrize('token,unit',[('NaN','OMR'),('1,3.2','OMR'),('-1','count'),('1.5','count')])
def test_numeric_corruption_rejected(token,unit):
    with pytest.raises(ValueError):
        decimal_value(token,unit)


def test_blank_source_value_is_not_zero():
    assert decimal_value('—','OMR','1000000') is None


def test_cumulative_period_cannot_be_invented():
    with pytest.raises(ValueError,match='explicit source period'):
        period({'year':'2025','frequency':'year_to_date'},{},date(2026,10,12))
    with pytest.raises(ValueError,match='future'):
        period({'year':'2026','frequency':'annual'},{},date(2026,10,12))


def test_empty_governorate_table_contains_no_observations():
    frame=actual_frame().iloc[:0].copy()
    frame['governorate']=pd.Series(dtype=str)
    config=table();config['geographic_level']='governorate'
    assert normalize_rows(frame,config,metadata(b''),Path('empty.csv'),'empty',date(2026,10,12))[0]==[]


def test_units_must_be_declared():
    config=table();config['metrics'][0].pop('source_unit')
    with pytest.raises(ValueError,match='source_unit'):
        normalize_rows(actual_frame(),config,metadata(b''),Path('cited.csv'),'test',date(2026,10,12))


def test_source_columns_cannot_be_overridden():
    config=table();config['constants']={'frequency':'annual'}
    with pytest.raises(ValueError,match='override'):
        normalize_rows(actual_frame(),config,metadata(b''),Path('cited.csv'),'test',date(2026,10,12))


def test_equal_genuine_duplicates_preserve_both_source_links():
    rows=normalized_actual_rows()
    alternate=deepcopy(rows[0]);alternate[1]['source_file']='second-cited-transcription.csv'
    observations,lineage,removed=deduplicate(rows+[alternate])
    assert len(observations)==1 and len(lineage)==2 and removed==1
    assert observations[0]['value']==genuine_national_observation()['value']


def test_conflicting_source_readings_block_deduplication():
    rows=normalized_actual_rows()
    alternate=deepcopy(rows[0])
    # Deliberately corrupt the value to a different genuinely published count;
    # this is a validation fault, not a created market observation.
    excerpt=(ROOT/'tests/fixtures/ncsi_2025_10_activity_excerpt.txt').read_text()
    source=json.loads((ROOT/'docs/source_catalog.json').read_text())[0]
    other=next(r['value'] for r in parse_page(excerpt,source,'test') if r['indicator']=='mortgage_contract_count')
    alternate[0]['value']=other
    with pytest.raises(ValueError,match='Conflicting'):
        deduplicate(rows+[alternate])


def test_national_records_do_not_establish_11_governorate_coverage():
    observations,_,_=deduplicate(normalized_actual_rows())
    report=coverage(observations)
    assert not report[0]['all_11_present']
    assert len(report[0]['missing_governorates'])==11


def test_empty_input_publishes_no_analytical_dataset(tmp_path):
    folder=tmp_path/'raw';folder.mkdir()
    output=tmp_path/'processed'
    report,code=run(folder,output,tmp_path/'status.json',date(2026,10,12))
    assert code==2 and report['unique_observations']==0
    assert not report['ready_for_python_powerbi'] and not output.exists()


def test_xlsx_empty_sheet_selection_and_original_preservation(tmp_path):
    # Header-only workbook: no synthetic or sample market records.
    path=tmp_path/'header-only.xlsx'
    book=Workbook();book.active.title='Data';book.active.append(['Year','Governorate','Transaction count']);book.save(path)
    before=path.read_bytes()
    config=table();config.update(sheet='Data',geographic_level='governorate')
    frame=read_table(path,config)
    assert frame.empty
    assert path.read_bytes()==before
    init_metadata(path)
    sidecar=Path(str(path)+'.metadata.json')
    assert json.loads(sidecar.read_text())['expected_sha256']==digest(before)
    with pytest.raises(ValueError,match='not be overwritten'):
        init_metadata(path)


def test_row_level_grain_is_rejected_before_reading(tmp_path):
    with pytest.raises(ValueError,match='aggregate'):
        read_table(tmp_path/'not-read.csv',{'observation_grain':'transaction'})


def test_missing_metadata_blocks_upload(tmp_path):
    path=tmp_path/'real-excerpt.csv'
    actual_frame().drop(columns='__source_row').to_csv(path,index=False)
    before=path.read_bytes()
    output=tmp_path/'out'
    report,code=run(tmp_path,output,tmp_path/'status.json',date(2026,10,12))
    assert code==1 and not output.exists()
    assert path.read_bytes()==before


def test_real_excerpt_roundtrip_is_partial_and_idempotent(tmp_path):
    folder=tmp_path/'inputs';folder.mkdir()
    path=folder/'cited-test-transcription.csv'
    actual_frame().drop(columns='__source_row').to_csv(path,index=False)
    before=path.read_bytes()
    Path(str(path)+'.metadata.json').write_text(json.dumps(metadata(before)))
    report,code=run(folder,tmp_path/'out',tmp_path/'status.json',date(2026,10,12))
    assert code==2 and report['unique_observations']==1 and not report['ready_for_python_powerbi']
    generated=Path(report['output_directory'])/'observations.csv'
    assert pd.read_csv(generated)['geographic_level'].tolist()==['national']
    original_output=generated.read_bytes()
    report2,code2=run(folder,tmp_path/'out',tmp_path/'status.json',date(2026,10,12))
    assert code2==2 and report2['output_directory']==report['output_directory']
    assert generated.read_bytes()==original_output and path.read_bytes()==before
    generated.write_text('technical corruption; not a market record')
    report3,code3=run(folder,tmp_path/'out',tmp_path/'status.json',date(2026,10,12))
    assert code3==1 and report3['state']=='validation_failed'
    assert generated.read_text()=='technical corruption; not a market record'


def test_checksum_and_review_admission():
    payload=b'header-only technical file'
    meta=metadata(payload)
    validate_metadata(meta,payload)
    with pytest.raises(ValueError,match='SHA-256'):
        validate_metadata(meta,payload+b'changed')
    meta['verification']['status']='pending'
    with pytest.raises(ValueError,match='reviewed'):
        validate_metadata(meta,payload)


def test_wide_year_workbook_uses_a_genuine_national_annual_count(tmp_path):
    source=json.loads((ROOT/'docs/source_catalog.json').read_text())[0]
    excerpt=(ROOT/'tests/fixtures/ncsi_2025_10_activity_excerpt.txt').read_text()
    actual=next(r for r in parse_page(excerpt,source,'test-transcription-only')
                if r['indicator']=='sales_contract_count' and r['frequency']=='annual')
    path=tmp_path/'cited-transcription-wide.xlsx'
    book=Workbook();book.active.title='Cited table';book.active.append([2024]);book.active.append([int(actual['value'])]);book.save(path)
    before=path.read_bytes()
    config=table();config.update(sheet='Cited table',layout='wide',wide_dimension='year',
        wide_columns=['2024'],wide_value_field='transaction_count',constants={'frequency':'annual'})
    rows,_,_=normalize_rows(read_table(path,config),config,metadata(before),path,digest(before),date(2026,10,12))
    assert rows[0][0]['value']==actual['value']
    assert rows[0][0]['period_start']=='2024-01-01'
    assert rows[0][1]['source_column']=='2024' and rows[0][1]['source_row']==2
    assert path.read_bytes()==before


def test_wide_governorate_header_only_is_reference_not_market_data(tmp_path):
    path=tmp_path/'header-only-wide.xlsx'
    book=Workbook();book.active.title='Headers';book.active.append([g[2] for g in GOVERNORATES]);book.save(path)
    config=table();config.update(sheet='Headers',layout='wide',wide_dimension='governorate',
        wide_columns=[g[2] for g in GOVERNORATES],wide_value_field='transaction_count')
    frame=read_table(path,config)
    assert frame.empty and set(frame.columns)=={'__source_row','__source_column','governorate','transaction_count'}


def test_arabic_frequency_on_genuine_reporting_interval():
    actual=genuine_national_observation()
    start,end,freq=period({'period_start':actual['period_start'],'period_end':actual['period_end'],'frequency':'شهري'}, {}, date(2026,10,12))
    assert (start,end,freq)==(actual['period_start'],actual['period_end'],'monthly')


def test_missing_real_metric_cell_is_logged_without_imputation():
    frame=actual_frame();frame['transaction_count']='—'
    rows,missing,_=normalize_rows(frame,table(),metadata(b''),Path('test-corrupted-missing-cell.csv'),
        'test-only',date(2026,10,12))
    assert rows==[] and len(missing)==1 and missing[0]['indicator']=='sales_contract_count'


def test_pending_licence_blocks_ingestion():
    meta=metadata(b'header-only technical bytes')
    meta['licence']['redistribution_allowed']=False
    with pytest.raises(ValueError,match='licence'):
        validate_metadata(meta,b'header-only technical bytes')


def test_index_definition_requires_base_period_without_creating_an_index():
    frame=actual_frame().iloc[:0].rename(columns={'transaction_count':'property_price_index'})
    config=table();config['metrics']=[dict(field='property_price_index',indicator='property_price_index',source_unit='index points',unit='index_points')]
    with pytest.raises(ValueError,match='base period'):
        normalize_rows(frame,config,metadata(b''),Path('empty.csv'),'empty',date(2026,10,12))


def test_presence_gate_uses_reference_labels_without_market_records():
    from src.manual_ingestion import governorate_presence
    labels={g[0] for g in GOVERNORATES}
    assert governorate_presence(labels)['all_11_present']
    missing=governorate_presence(labels-{'muscat'})
    assert not missing['all_11_present'] and missing['missing_governorates']==['muscat']
    assert not governorate_presence(labels|{'national_total'})['all_11_present']
