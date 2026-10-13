"""Read XLSX worksheet XML independently of unreliable declared dimensions."""
from collections import Counter
from pathlib import Path
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def read_workbook(path):
    """Return actual rows, source XML number tokens and diagnostics; never save."""
    path = Path(path)
    with ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {item.attrib['Id']:item.attrib['Target'] for item in relationships}
        xml_sheets = {}
        for item in workbook.find('s:sheets',NS):
            target = targets[item.attrib['{'+NS['r']+'}id']]
            target = target.lstrip('/') if target.startswith('/') else 'xl/'+target
            element = ET.fromstring(archive.read(target))
            dimension = element.find('s:dimension',NS)
            xml_sheets[item.attrib['name']] = dict(
                declared_xml_range=dimension.attrib.get('ref') if dimension is not None else None,
                xml_cells=len(element.findall('.//s:sheetData/s:row/s:c',NS)),
                formula_cells=len(element.findall('.//s:f',NS)),
                numeric_tokens={cell.attrib['r']:cell.find('s:v',NS).text
                    for cell in element.findall('.//s:sheetData/s:row/s:c',NS)
                    if cell.attrib.get('t','n')=='n' and cell.find('s:v',NS) is not None},
            )
        core = ET.fromstring(archive.read('docProps/core.xml')) if 'docProps/core.xml' in archive.namelist() else None
        core_properties = {item.tag.split('}')[-1]:item.text for item in core} if core is not None else {}
    with warnings.catch_warnings(record=True) as observed:
        warnings.simplefilter('always')
        book = load_workbook(path,read_only=True,data_only=True)
        sheets = []
        try:
            for sheet in book:
                declared = dict(rows=sheet.max_row,columns=sheet.max_column)
                # This repairs iteration bounds in memory, leaving ZIP bytes intact.
                sheet.reset_dimensions()
                rows = list(sheet.iter_rows(values_only=True))
                actual = dict(rows=len(rows),columns=max((len(row) for row in rows),default=0))
                sheets.append(dict(name=sheet.title,rows=rows,declared_dimensions=declared,
                    actual_dimensions=actual,dimension_reset_required=declared!=actual,**xml_sheets[sheet.title]))
        finally:
            book.close()
    return dict(sheets=sheets,core_properties=core_properties,
                reader_warnings=sorted({str(item.message) for item in observed}))


def column_types(values):
    return dict(Counter('missing' if value is None else type(value).__name__ for value in values))
